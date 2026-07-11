from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path


class DocumentParseError(ValueError):
    pass


class UnsupportedDocumentError(DocumentParseError):
    pass


class EmptyDocumentError(DocumentParseError):
    pass


class OcrRequiredError(DocumentParseError):
    pass


@dataclass(frozen=True)
class LocatedText:
    content: str
    location_label: str
    heading: str = ""


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    blocks: list[LocatedText]


@dataclass(frozen=True)
class DocumentChunkInput:
    position: int
    location_label: str
    heading: str
    content: str


EXTENSION_TYPES = {
    ".md": "markdown",
    ".markdown": "markdown",
    ".txt": "text",
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".png": "image",
    ".jpg": "image",
    ".jpeg": "image",
    ".webp": "image",
}

MIME_TYPES = {
    "text/markdown": "markdown",
    "text/plain": "text",
    "application/pdf": "pdf",
    "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
    "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
    "image/png": "image",
    "image/jpeg": "image",
    "image/webp": "image",
}


def detect_document_type(filename: str, mime_type: str) -> str:
    document_type = EXTENSION_TYPES.get(Path(filename).suffix.lower())
    if document_type is None:
        document_type = MIME_TYPES.get(mime_type.lower().split(";", 1)[0].strip())
    if document_type is None:
        raise UnsupportedDocumentError(f"Unsupported document type: {filename}")
    return document_type


def parse_document(path: Path, document_type: str) -> ParsedDocument:
    parser = {
        "markdown": _parse_markdown,
        "text": _parse_text,
        "pdf": _parse_pdf,
        "docx": _parse_docx,
        "pptx": _parse_pptx,
    }.get(document_type)
    if document_type == "image":
        raise OcrRequiredError("OCR is required before this image can be indexed")
    if parser is None:
        raise UnsupportedDocumentError(f"Unsupported document type: {document_type}")
    blocks = parser(path)
    if not blocks:
        raise EmptyDocumentError(f"Document is empty or has no extractable text: {path.name}")
    return ParsedDocument(title=path.stem, blocks=blocks)


def chunk_located_text(
    blocks: list[LocatedText],
    max_chars: int = 1800,
    overlap_chars: int = 180,
) -> list[DocumentChunkInput]:
    if max_chars < 1:
        raise ValueError("max_chars must be positive")
    if overlap_chars < 0 or overlap_chars >= max_chars:
        raise ValueError("overlap_chars must be between 0 and max_chars - 1")

    chunks: list[DocumentChunkInput] = []
    step = max_chars - overlap_chars
    for block in blocks:
        content = block.content.strip()
        if not content:
            continue
        for start in range(0, len(content), step):
            piece = content[start : start + max_chars]
            if not piece:
                break
            chunks.append(
                DocumentChunkInput(
                    position=len(chunks),
                    location_label=block.location_label,
                    heading=block.heading,
                    content=piece,
                )
            )
            if start + max_chars >= len(content):
                break
    return chunks


def _normalize(value: str) -> str:
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _parse_markdown(path: Path) -> list[LocatedText]:
    blocks: list[LocatedText] = []
    heading = "document"
    content_lines: list[str] = []

    def flush() -> None:
        content = _normalize("\n".join(content_lines))
        if content:
            blocks.append(LocatedText(content, heading, heading))
        content_lines.clear()

    for line in path.read_text(encoding="utf-8").splitlines():
        match = re.match(r"^#{1,6}\s+(.+?)\s*$", line)
        if match:
            flush()
            heading = match.group(1).strip()
        else:
            content_lines.append(line)
    flush()
    return blocks


def _parse_text(path: Path) -> list[LocatedText]:
    content = _normalize(path.read_text(encoding="utf-8"))
    return [LocatedText(content, "text")] if content else []


def _parse_pdf(path: Path) -> list[LocatedText]:
    from pypdf import PdfReader

    blocks: list[LocatedText] = []
    for index, page in enumerate(PdfReader(str(path)).pages, start=1):
        content = _normalize(page.extract_text() or "")
        if content:
            blocks.append(LocatedText(content, f"page {index}"))
    return blocks


def _parse_docx(path: Path) -> list[LocatedText]:
    from docx import Document

    blocks: list[LocatedText] = []
    heading = ""
    for index, paragraph in enumerate(Document(str(path)).paragraphs, start=1):
        content = _normalize(paragraph.text)
        if not content:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            heading = content
            continue
        blocks.append(LocatedText(content, f"paragraph {index}", heading))
    return blocks


def _parse_pptx(path: Path) -> list[LocatedText]:
    from pptx import Presentation

    blocks: list[LocatedText] = []
    for index, slide in enumerate(Presentation(str(path)).slides, start=1):
        texts = [
            _normalize(shape.text)
            for shape in slide.shapes
            if hasattr(shape, "text") and _normalize(shape.text)
        ]
        if texts:
            blocks.append(
                LocatedText(
                    content="\n".join(texts),
                    location_label=f"slide {index}",
                    heading=texts[0],
                )
            )
    return blocks
