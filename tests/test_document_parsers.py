from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation

from app.services.documents import (
    EmptyDocumentError,
    LocatedText,
    OcrRequiredError,
    UnsupportedDocumentError,
    chunk_located_text,
    detect_document_type,
    parse_document,
)


def write_minimal_pdf(path: Path, text: str) -> None:
    escaped = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    stream = f"BT /F1 16 Tf 72 720 Td ({escaped}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        (
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>"
        ),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
        + stream
        + b"\nendstream",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode("ascii"))
        payload.extend(obj)
        payload.extend(b"\nendobj\n")
    xref_offset = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n".encode("ascii"))
    payload.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        payload.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    payload.extend(
        (
            f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
            f"startxref\n{xref_offset}\n%%EOF\n"
        ).encode("ascii")
    )
    path.write_bytes(payload)


def test_markdown_parser_preserves_heading_location(tmp_path):
    path = tmp_path / "lecture.md"
    path.write_text("# Entropy\nEntropy measures uncertainty.\n", encoding="utf-8")

    parsed = parse_document(path, detect_document_type(path.name, "text/markdown"))

    assert parsed.title == "lecture"
    assert parsed.blocks[0].location_label == "Entropy"
    assert "uncertainty" in parsed.blocks[0].content


def test_plain_text_parser_uses_text_location(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("Plain text study note.", encoding="utf-8")

    parsed = parse_document(path, detect_document_type(path.name, "text/plain"))

    assert parsed.blocks[0].location_label == "text"
    assert parsed.blocks[0].content == "Plain text study note."


def test_pdf_parser_preserves_page_location(tmp_path):
    path = tmp_path / "entropy.pdf"
    write_minimal_pdf(path, "Entropy measures uncertainty")

    parsed = parse_document(path, detect_document_type(path.name, "application/pdf"))

    assert parsed.blocks[0].location_label == "page 1"
    assert "Entropy measures uncertainty" in parsed.blocks[0].content


def test_docx_parser_preserves_paragraph_location(tmp_path):
    path = tmp_path / "lecture.docx"
    document = Document()
    document.add_heading("Channel Coding", level=1)
    document.add_paragraph("A document paragraph about redundancy.")
    document.save(path)

    parsed = parse_document(
        path,
        detect_document_type(
            path.name,
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ),
    )

    assert any(block.location_label == "paragraph 2" for block in parsed.blocks)
    assert any("redundancy" in block.content for block in parsed.blocks)


def test_pptx_parser_preserves_slide_location(tmp_path):
    path = tmp_path / "lecture.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Routing"
    slide.placeholders[1].text = "A presentation slide about shortest paths."
    presentation.save(path)

    parsed = parse_document(
        path,
        detect_document_type(
            path.name,
            "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        ),
    )

    assert parsed.blocks[0].location_label == "slide 1"
    assert "shortest paths" in parsed.blocks[0].content


def test_image_requires_ocr_provider(tmp_path):
    path = tmp_path / "board.png"
    path.write_bytes(b"\x89PNG\r\n\x1a\n")

    assert detect_document_type(path.name, "image/png") == "image"
    with pytest.raises(OcrRequiredError, match="OCR"):
        parse_document(path, "image")


def test_unsupported_extension_is_rejected():
    with pytest.raises(UnsupportedDocumentError, match="Unsupported"):
        detect_document_type("archive.zip", "application/zip")


def test_empty_text_document_is_rejected(tmp_path):
    path = tmp_path / "empty.txt"
    path.write_text("   \n", encoding="utf-8")

    with pytest.raises(EmptyDocumentError, match="empty"):
        parse_document(path, "text")


def test_chunking_keeps_location_and_overlap():
    blocks = [LocatedText(content="A" * 1200, location_label="page 2", heading="Entropy")]

    chunks = chunk_located_text(blocks, max_chars=500, overlap_chars=50)

    assert len(chunks) == 3
    assert {chunk.location_label for chunk in chunks} == {"page 2"}
    assert chunks[0].content[-50:] == chunks[1].content[:50]
